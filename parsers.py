import os
from pathlib import Path
from typing import Optional
import docx
import pypdf

def extract_text_from_file(file_path: str | Path) -> str:
    """Extrai texto de arquivos .txt, .md, .docx, .pdf, .csv, .json"""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")

    suffix = path.suffix.lower()

    if suffix in [".txt", ".md", ".csv", ".json", ".log"]:
        return extract_text_from_plain(path)
    elif suffix in [".docx", ".doc"]:
        return extract_text_from_docx(path)
    elif suffix == ".pdf":
        return extract_text_from_pdf(path)
    else:
        # Tenta ler como texto genérico
        try:
            return extract_text_from_plain(path)
        except Exception:
            raise ValueError(f"Formato de arquivo '{suffix}' não suportado diretamente para extração de texto.")

def extract_text_from_plain(path: Path) -> str:
    encodings = ["utf-8", "latin-1", "cp1252", "iso-8859-1"]
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Não foi possível decodificar o arquivo de texto {path.name}")

def extract_text_from_docx(path: Path) -> str:
    doc = docx.Document(str(path))
    full_text = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)

    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            table_rows.append(" | ".join(row_data))
        if table_rows:
            full_text.append("\n--- Tabela ---\n" + "\n".join(table_rows) + "\n--------------\n")

    return "\n\n".join(full_text)

def extract_text_from_pdf(path: Path) -> str:
    reader = pypdf.PdfReader(str(path))
    pages_text = []
    for idx, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            pages_text.append(f"--- Página {idx + 1} ---\n" + text.strip())

    if not pages_text:
        return "[Aviso: O PDF pode ser baseado em imagem escaneada ou não contém texto selecionável]"

    return "\n\n".join(pages_text)
