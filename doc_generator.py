import re
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
from datetime import datetime

# DOCX
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# ReportLab PDF
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether, Image as RLImage
from reportlab.pdfgen import canvas

PRIMARY_HEX = "#1A365D"    # Deep Navy
SECONDARY_HEX = "#2B6CB0"  # Slate Blue
TEXT_DARK_HEX = "#2D3748"  # Charcoal
BG_LIGHT_HEX = "#F7FAFC"   # Very Light Gray
BORDER_HEX = "#E2E8F0"     # Soft Gray

PRIMARY_COLOR = RGBColor(0x1A, 0x36, 0x5D)
SECONDARY_COLOR = RGBColor(0x2B, 0x6C, 0xB0)
TEXT_DARK = RGBColor(0x2D, 0x37, 0x48)

class NumberedCanvas(canvas.Canvas):
    """Canvas de duas passadas para inserir 'Página X de Y' no rodapé do PDF"""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#718096"))
        
        # Linha separadora do rodapé
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 45, 595 - 54, 45)

        footer_text = f"Página {self._pageNumber} de {page_count}"
        self.drawRightString(595 - 54, 32, footer_text)
        self.drawString(54, 32, "Gerado por Assistente de Relatórios")
        self.restoreState()

def parse_markdown_blocks(markdown_text: str) -> List[Dict[str, Any]]:
    """
    Analisa o Markdown em blocos semânticos:
    - heading (level, text)
    - paragraph (text)
    - list_item (ordered, text)
    - table (headers, rows)
    """
    lines = markdown_text.strip().split("\n")
    blocks = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        # Títulos
        if line.startswith("#"):
            level = 0
            while level < len(line) and line[level] == "#":
                level += 1
            text = line[level:].strip()
            blocks.append({"type": "heading", "level": min(level, 4), "text": text})
            i += 1
            continue

        # Tabelas Markdown (| col1 | col2 | ...)
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                # Processa linhas da tabela
                rows = []
                for tline in table_lines:
                    # Ignora linhas divisórias como |---|---|
                    cells = [c.strip() for c in tline.strip("|").split("|")]
                    if all(re.match(r"^:?-+:?$", c) for c in cells if c):
                        continue
                    rows.append(cells)

                if rows:
                    headers = rows[0]
                    data_rows = rows[1:]
                    blocks.append({"type": "table", "headers": headers, "rows": data_rows})
            continue

        # Listas com marcadores (- ou *)
        if line.strip().startswith(("- ", "* ", "• ")):
            item_text = line.strip()[2:].strip()
            blocks.append({"type": "list_item", "ordered": False, "text": item_text})
            i += 1
            continue

        # Listas numeradas (1. , 2. )
        num_match = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if num_match:
            item_text = num_match.group(2).strip()
            blocks.append({"type": "list_item", "ordered": True, "text": item_text})
            i += 1
            continue

        # Parágrafo comum (acumula linhas subsequentes se houver)
        para_lines = [line]
        i += 1
        while i < n and lines[i].strip() and not lines[i].startswith(("#", "|", "- ", "* ", "• ")) and not re.match(r"^\s*\d+\.\s+", lines[i]):
            para_lines.append(lines[i].strip())
            i += 1

        blocks.append({"type": "paragraph", "text": " ".join(para_lines)})

    return blocks

def _apply_cell_shading(cell, color_hex: str):
    """Aplica cor de fundo a uma célula de tabela no Word"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex.lstrip("#")}"/>')
    tcPr.append(shd)

def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Define margens internas de célula de tabela no Word"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_docx_report(markdown_text: str, output_path: str | Path, title_hint: str = "Relatório", image_paths: Optional[List[Path]] = None) -> Path:
    """Gera um documento Microsoft Word (.docx) estilizado a partir do Markdown"""
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # Configuração de margens (2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    blocks = parse_markdown_blocks(markdown_text)

    first_h1 = True

    for block in blocks:
        btype = block["type"]

        if btype == "heading":
            level = block["level"]
            text = block["text"]

            if level == 1 and first_h1:
                # Título Principal do Documento
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(text)
                run.font.name = "Calibri"
                run.font.size = Pt(24)
                run.font.bold = True
                run.font.color.rgb = PRIMARY_COLOR

                # Subtítulo com data
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.space_before = Pt(0)
                p_sub.paragraph_format.space_after = Pt(16)
                date_str = datetime.now().strftime("%d de %B de %Y")
                run_sub = p_sub.add_run(f"Data de Emissão: {date_str} • Status: Oficial")
                run_sub.font.name = "Calibri"
                run_sub.font.size = Pt(10)
                run_sub.font.italic = True
                run_sub.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

                first_h1 = False
                continue

            # Seções e Subseções
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            p.paragraph_format.space_after = Pt(4)

            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.bold = True

            if level == 1:
                run.font.size = Pt(16)
                run.font.color.rgb = PRIMARY_COLOR
            elif level == 2:
                run.font.size = Pt(13)
                run.font.color.rgb = SECONDARY_COLOR
            else:
                run.font.size = Pt(11)
                run.font.color.rgb = TEXT_DARK

        elif btype == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            # Formata negritos simples (**texto**)
            parts = re.split(r"(\*\*.*?\*\*)", block["text"])
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = "Calibri"
                run.font.size = Pt(11)
                run.font.color.rgb = TEXT_DARK

        elif btype == "list_item":
            p = doc.add_paragraph(style="List Bullet" if not block["ordered"] else "List Number")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r"(\*\*.*?\*\*)", block["text"])
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = "Calibri"
                run.font.size = Pt(10.5)
                run.font.color.rgb = TEXT_DARK

        elif btype == "table":
            headers = block["headers"]
            rows = block["rows"]
            cols_count = max(len(headers), max((len(r) for r in rows), default=0))

            if cols_count == 0:
                continue

            table = doc.add_table(rows=1 + len(rows), cols=cols_count)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Cabeçalho
            hdr_cells = table.rows[0].cells
            for col_idx in range(cols_count):
                cell = hdr_cells[col_idx]
                val = headers[col_idx] if col_idx < len(headers) else ""
                cell.text = val
                _apply_cell_shading(cell, PRIMARY_HEX)
                _set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for cp in cell.paragraphs:
                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in cp.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(10)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            # Linhas de dados
            for r_idx, row_data in enumerate(rows):
                row_cells = table.rows[r_idx + 1].cells
                bg_color = BG_LIGHT_HEX if r_idx % 2 == 1 else "#FFFFFF"
                for col_idx in range(cols_count):
                    cell = row_cells[col_idx]
                    val = row_data[col_idx] if col_idx < len(row_data) else ""
                    cell.text = val
                    _apply_cell_shading(cell, bg_color)
                    _set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    for cp in cell.paragraphs:
                        for run in cp.runs:
                            run.font.name = "Calibri"
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = TEXT_DARK

            # Espaçamento após tabela
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_before = Pt(6)
            p_space.paragraph_format.space_after = Pt(6)

    # Anexa as imagens/fotos analisadas ao final do documento
    if image_paths:
        valid_images = [Path(p) for p in image_paths if Path(p).exists()]
        if valid_images:
            p_sec = doc.add_paragraph()
            p_sec.paragraph_format.space_before = Pt(16)
            p_sec.paragraph_format.space_after = Pt(6)
            run_sec = p_sec.add_run("📸 Imagens / Anexos Analisados")
            run_sec.font.name = "Calibri"
            run_sec.font.size = Pt(13)
            run_sec.font.bold = True
            run_sec.font.color.rgb = SECONDARY_COLOR

            for img_p in valid_images:
                try:
                    doc.add_picture(str(img_p), width=Inches(5.0))
                    p_cap = doc.add_paragraph()
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(10)
                    r_cap = p_cap.add_run(f"Anexo: {img_p.name}")
                    r_cap.font.name = "Calibri"
                    r_cap.font.size = Pt(9)
                    r_cap.font.italic = True
                    r_cap.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
                except Exception:
                    pass

    doc.save(str(out_file))
    return out_file

def _clean_md_for_pdf(text: str) -> str:
    """Converte Markdown inline (**negrito**, *itálico*) para tags HTML aceitas pelo ReportLab"""
    # Escapa caracteres XML básicos
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    # Restaura tags seguras
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.*?)\*", r"<i>\1</i>", text)
    return text

def create_pdf_report(markdown_text: str, output_path: str | Path, title_hint: str = "Relatório", image_paths: Optional[List[Path]] = None) -> Path:
    """Gera um documento PDF com diagramação e tipografia executiva a partir do Markdown"""
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(out_file),
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        textColor=colors.HexColor(PRIMARY_HEX),
        spaceAfter=4
    )

    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#718096"),
        spaceAfter=14
    )

    h1_style = ParagraphStyle(
        "ReportH1",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        textColor=colors.HexColor(PRIMARY_HEX),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        "ReportH2",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor(SECONDARY_HEX),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        "ReportBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor(TEXT_DARK_HEX),
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        "ReportBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor(TEXT_DARK_HEX),
        leftIndent=14,
        spaceAfter=3
    )

    table_cell_style = ParagraphStyle(
        "ReportTableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor(TEXT_DARK_HEX)
    )

    table_hdr_style = ParagraphStyle(
        "ReportTableHdr",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white
    )

    story = []
    blocks = parse_markdown_blocks(markdown_text)
    first_h1 = True

    usable_width = 595.27 - 108  # A4 width - margins

    for block in blocks:
        btype = block["type"]

        if btype == "heading":
            level = block["level"]
            raw_text = block["text"]

            if level == 1 and first_h1:
                story.append(Paragraph(_clean_md_for_pdf(raw_text), title_style))
                date_str = datetime.now().strftime("%d de %B de %Y")
                story.append(Paragraph(f"Data de Emissão: {date_str} • Status: Oficial", subtitle_style))
                story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor(PRIMARY_HEX), spaceAfter=14))
                first_h1 = False
                continue

            if level == 1:
                story.append(Paragraph(_clean_md_for_pdf(raw_text), h1_style))
            else:
                story.append(Paragraph(_clean_md_for_pdf(raw_text), h2_style))

        elif btype == "paragraph":
            story.append(Paragraph(_clean_md_for_pdf(block["text"]), body_style))

        elif btype == "list_item":
            bullet_char = "•" if not block["ordered"] else "&bull;"
            formatted_text = f"{bullet_char} {_clean_md_for_pdf(block['text'])}"
            story.append(Paragraph(formatted_text, bullet_style))

        elif btype == "table":
            headers = block["headers"]
            rows = block["rows"]
            cols_count = max(len(headers), max((len(r) for r in rows), default=0))

            if cols_count == 0:
                continue

            col_width = usable_width / cols_count
            table_data = []

            # Cabeçalho da tabela
            hdr_cells = []
            for col_idx in range(cols_count):
                val = headers[col_idx] if col_idx < len(headers) else ""
                hdr_cells.append(Paragraph(_clean_md_for_pdf(val), table_hdr_style))
            table_data.append(hdr_cells)

            # Linhas
            for r_idx, row_data in enumerate(rows):
                row_cells = []
                for col_idx in range(cols_count):
                    val = row_data[col_idx] if col_idx < len(row_data) else ""
                    row_cells.append(Paragraph(_clean_md_for_pdf(val), table_cell_style))
                table_data.append(row_cells)

            t_style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(PRIMARY_HEX)),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(BORDER_HEX)),
            ]

            # Linhas com zebrado suave
            for r_idx in range(1, len(table_data)):
                if r_idx % 2 == 0:
                    t_style.append(("BACKGROUND", (0, r_idx), (-1, r_idx), colors.HexColor(BG_LIGHT_HEX)))

            pdf_table = Table(table_data, colWidths=[col_width] * cols_count)
            pdf_table.setStyle(TableStyle(t_style))
            story.append(Spacer(1, 6))
            story.append(pdf_table)
            story.append(Spacer(1, 10))

    # Anexa as fotos/imagens analisadas no PDF
    if image_paths:
        valid_images = [Path(p) for p in image_paths if Path(p).exists()]
        if valid_images:
            story.append(Spacer(1, 14))
            story.append(Paragraph("<b>📸 Imagens / Anexos Analisados</b>", h2_style))
            story.append(Spacer(1, 8))
            for img_p in valid_images:
                try:
                    story.append(RLImage(str(img_p), width=450, height=260, kind='proportional'))
                    story.append(Spacer(1, 10))
                except Exception:
                    pass

    doc.build(story, canvasmaker=NumberedCanvas)
    return out_file

def create_txt_report(markdown_text: str, output_path: str | Path) -> Path:
    """Gera o arquivo em texto puro / Markdown"""
    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(markdown_text)
    return out_file
